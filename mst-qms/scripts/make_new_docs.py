"""신규 제정 지침서·양식을 MST 원본 문서 틀 그대로 만든다.
 - 지침서: MST 원본 지침서(MI-0804 치공구 관리 지침서) xlsx를 틀로 사용 (표지·개정이력 / 본문 / 프로세스 맵 3개 시트, 로고 유지)
 - 양식  : MST 원본 양식(MI-0804-001 치공구 관리 대장) xlsx를 틀로 사용 (로고·결재란 유지)
사용: python3 make_new_docs.py <지침서틀.xlsx> <양식틀.xlsx> <출력 폴더>
"""
import copy, os, sys, zipfile, io
import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Side, Font
from openpyxl.utils import get_column_letter

TPL_DOC, TPL_FORM, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
os.makedirs(OUT, exist_ok=True)
DATE = '2026-10-05'

def logo_bytes(path):
    z = zipfile.ZipFile(path)
    names = sorted([n for n in z.namelist() if n.startswith('xl/media/')], key=lambda n: z.getinfo(n).file_size)
    return z.read(names[0])  # 가장 작은 이미지 = MST 로고

def add_logo(ws, data, anchor='A1', w=102, h=30):
    img = Image(io.BytesIO(data)); img.width, img.height = w, h; ws.add_image(img, anchor)

def wrap_lines(text, width=46, indent='    '):
    """MST 본문처럼 한 줄 약 46자로 나누어 여러 셀에 기입"""
    out, line = [], ''
    for word in text.split(' '):
        if len(line) + len(word) + 1 > width and line:
            out.append(line); line = indent + word
        else:
            line = (line + ' ' + word) if line else word
    if line: out.append(line)
    return out

# ───────── 지침서 ─────────
DOCS = [
  dict(code='MI-0703', title='클린룸·검사실 환경 관리 지침서', proc='생산관리', team='품질팀', map_title='클린룸·검사실 환경 관리 프로세스 맵',
    body=[
      ('1. 목적', ["본 지침서는 주식회사엠에스티(이하 '회사'라 한다)에서 생산하는 반도체 장비 부품(Chuck)의 오염을 방지하기 위하여 클린룸(준클린룸 포함)과 수입·출하 검사실의 청정도, 온·습도, 작업자 복장 및 정리·정돈 상태를 관리하는 기준을 정하는 데 그 목적이 있다."]),
      ('2. 적용범위', ["회사의 클린룸(준클린룸 포함), 수입검사실, 출하검사실 및 포장실의 환경 관리 업무에 적용한다."]),
      ('3. 용어의 정의', ['3.1 클린룸', "    공기 중 부유 입자(Particle)를 관리 기준 이하로 유지하도록 설계·운영되는 작업 구역을 말한다.",
                       '3.2 준클린룸', "    클린룸 수준의 설비는 없으나 출입 통제, 복장, 청소 기준을 적용하여 관리하는 구역을 말한다.",
                       '3.3 Particle Spec', "    측정 지점별로 정한 단위 체적당 허용 입자 수 기준을 말한다."]),
      ('4. 책임과 권한', ['4.1 품질팀장', "    환경 관리 기준의 제정·개정, Particle 및 온·습도 측정 결과 검토, 기준 이탈 시 조치를 지시한다.",
                        '4.2 제조팀장', "    클린룸 5S 점검과 작업자 복장 점검을 실시하고, 기준 이탈 사항을 품질팀에 통보한다.",
                        '4.3 관리팀장', "    클린복·클린화 등 보호구의 구매, 세탁 및 교체를 지원한다."]),
      ('5. 관리 기준', ['5.1 구역 지정', "    클린룸, 준클린룸, 검사실 구역을 Lay-out에 표시하고 구역별 출입 기준을 출입구에 게시한다.",
                      '5.2 청정도 (Particle)', "    품질팀은 지정된 측정 지점에서 월 1회 Particle을 측정하여 Particle 측정 관리대장(MI-0703-002)에 기록한다. Spec 초과 시 원인을 조사하고 개선 후 재측정한다.",
                      '5.3 온·습도', "    수입·출하 검사실은 온도 20±2℃, 습도 65% 이하로 유지하며, 매일 09:00와 15:00에 검사실 온·습도 관리표(MI-0703-004)에 기록한다.",
                      '5.4 클린룸 5S', "    제조팀은 매일 작업 시작 전 정리·정돈·청소·청결·습관화 항목을 클린룸 5S 점검표(MI-0703-001)로 점검한다.",
                      '5.5 복장 관리', "    클린룸 출입자는 클린복, 클린화, 모자, 마스크, 장갑을 착용한다. 제조팀은 매주 착용·세탁·교체 상태를 클린복·보호구 관리 점검표(MI-0703-003)로 점검한다.",
                      '5.6 검사품 구분', "    검사실 내 검사품은 대기품, 합격품, 불합격품으로 구분하여 지정된 위치에 보관하고 식별표를 부착한다."]),
      ('6. 이상 발생 시 조치', ["기준을 벗어난 경우 해당 구역의 작업을 중지하고, 부적합 및 시정조치 절차서(MD-1002)에 따라 원인 분석과 재발 방지 대책을 수립한다. 이미 생산된 제품은 영향 여부를 확인한다."]),
      ('7. 관련 양식', ['MI-0703-001 클린룸 5S 점검표', 'MI-0703-002 Particle 측정 관리대장', 'MI-0703-003 클린복·보호구 관리 점검표', 'MI-0703-004 검사실 온·습도 관리표']),
      ('8. 기록 보존', ["본 지침서에 따른 기록은 문서화된 정보 관리 절차서(MD-0702)에 따라 3년간 보존한다."]),
    ],
    flow=[('구역 지정 기준', '구역 지정·출입 기준 게시', 'Lay-out, 출입 기준', '품질팀', ''),
          ('측정 계획', 'Particle 측정 (월 1회)', 'MI-0703-002', '품질팀', 'Particle Spec 준수율'),
          ('검사실 환경', '온·습도 측정 (일 2회)', 'MI-0703-004', '품질팀', '온·습도 이탈 건수'),
          ('점검 기준', '클린룸 5S 점검 (매일)', 'MI-0703-001', '제조팀', ''),
          ('복장 기준', '클린복·보호구 점검 (매주)', 'MI-0703-003', '제조팀', ''),
          ('기준 이탈', '이상 조치 (MD-1002)', '시정조치 보고서', '품질팀', '')]),
  dict(code='MI-0810', title='정보보안 및 고객 도면 관리 지침서', proc='개발관리', team='개발팀', map_title='정보보안 및 고객 도면 관리 프로세스 맵',
    body=[
      ('1. 목적', ["본 지침서는 주식회사엠에스티(이하 '회사'라 한다)가 고객으로부터 받은 도면, 사양서 등 고객 재산(정보)과 회사의 기술 정보가 유출·훼손·오용되지 않도록 관리하는 방법을 정하는 데 그 목적이 있다."]),
      ('2. 적용범위', ["고객 도면, 사양서, 검사 기준 등 고객이 제공한 기술 정보와 이를 담은 전자 파일 및 출력물의 접수, 배포, 사용, 보관, 회수, 폐기에 적용한다."]),
      ('3. 용어의 정의', ['3.1 고객 도면', "    고객이 제품 제작을 위하여 회사에 제공한 도면 및 사양서를 말한다.",
                       '3.2 관리본', "    관리 번호와 배포처가 등록되어 개정 시 회수 대상이 되는 도면을 말한다."]),
      ('4. 책임과 권한', ['4.1 개발팀장', "    고객 도면의 접수, 등록, 배포, 회수 및 구본 폐기를 주관한다.",
                        '4.2 각 부서장', "    배포받은 도면의 사용과 보관 상태를 관리하며, 작업 종료 후 지정 보관함에 보관하도록 한다.",
                        '4.3 관리팀장', "    월 1회 정보보안 점검을 실시하고, 보안 교육을 주관한다."]),
      ('5. 업무 절차', ['5.1 도면 접수·등록', "    고객 도면은 접수 즉시 고객 도면 배포·회수 대장(MI-0810-001)에 등록하고 관리 번호를 부여한다.",
                      '5.2 배포', "    필요한 부서에만 관리본을 배포하고 배포 부서와 일자를 대장에 기록한다.",
                      '5.3 사용', "    도면은 작업 기준으로만 사용하며, 도면을 검사 성적서 대용으로 사용하지 않는다. 현장에 도면을 방치하지 않는다.",
                      '5.4 개정·회수', "    도면이 개정되면 구본을 회수하여 폐기하거나 '구본' 표시 후 별도 보관한다.",
                      '5.5 전자 파일', "    고객 도면 파일은 지정된 저장 위치에만 보관하며, 이동식 저장 매체와 개인 메일을 통한 반출을 금지한다.",
                      '5.6 보안 점검', "    관리팀은 월 1회 정보보안·도면 관리 점검표(MI-0810-002)로 현장 도면 방치, 구본 회수, 파일 반출 통제 상태를 점검한다."]),
      ('6. 위반 시 조치', ["점검 결과 위반 사항은 부적합 및 시정조치 절차서(MD-1002)에 따라 조치하고, 관련 인원에게 재교육을 실시한다."]),
      ('7. 관련 양식', ['MI-0810-001 고객 도면 배포·회수 대장', 'MI-0810-002 정보보안·도면 관리 점검표']),
      ('8. 기록 보존', ["본 지침서에 따른 기록은 문서화된 정보 관리 절차서(MD-0702)에 따라 3년간 보존한다."]),
    ],
    flow=[('고객 도면 입수', '도면 접수·등록', 'MI-0810-001', '개발팀', ''),
          ('작업 필요 부서', '관리본 배포', 'MI-0810-001', '개발팀', ''),
          ('작업 지시', '현장 사용·보관', '지정 보관함', '각 부서', ''),
          ('도면 개정', '구본 회수·폐기', 'MI-0810-001', '개발팀', '구본 회수율'),
          ('점검 계획', '보안 점검 (월 1회)', 'MI-0810-002', '관리팀', '보안 위반 건수'),
          ('위반 사항', '시정조치 (MD-1002)', '시정조치 보고서', '관리팀', '')]),
]

def make_doc(d, logo):
    wb = openpyxl.load_workbook(TPL_DOC)
    s1, s2 = wb.worksheets[0], wb.worksheets[1]
    s1['I1'] = d['proc'] + '  프 로 세 스'; s1['I2'] = d['title']; s1['AJ2'] = d['code']
    s1['A5'] = d['team']; s1['L5'] = d['team']; s1['X5'] = DATE
    for c in ('AJ5', 'AM5', 'AP5'): s1[c] = None
    s1['A8'] = '0.0'; s1['D8'] = DATE; s1['K8'] = 'ISO 9001:2015 · 고객(세메스) 요구사항 반영 신규 제정'; s1['AI8'] = '전면'
    if not s1._images: add_logo(s1, logo)
    s2['A1'] = d['title']; s2['F3'] = DATE.replace('-', '.'); s2['AB3'] = DATE.replace('-', '.'); s2['Q3'] = '0.0'
    # 본문 영역 비우기 (B5 이후 B·C열 텍스트)
    for row in s2.iter_rows(min_row=5, max_row=s2.max_row):
        for c in row:
            if c.__class__.__name__ != 'MergedCell' and c.column in (2, 3) and isinstance(c.value, str): c.value = None
    r = 5
    for head, paras in d['body']:
        s2.cell(r, 2, head); r += 1
        for p in paras:
            if p.startswith('    ') or not (p[:2].replace('.', '').isdigit() or p[:3].replace('.', '').isdigit()):
                lines = wrap_lines(p.strip() if not p.startswith('    ') else p.strip(), 46, '    ')
                for i, ln in enumerate(lines):
                    s2.cell(r, 2, ('    ' if p.startswith('    ') or i else '  ') + ln.strip()); r += 1
            else:
                s2.cell(r, 2, p); r += 1
        r += 1
    # 프로세스 맵: 표 형식 (입력 / Process / 출력 / 책임부문 / 성과지표)
    if len(wb.worksheets) > 2:
        s3 = wb.worksheets[2]
        s3.title = d['map_title'][:31]
        for mr in list(s3.merged_cells.ranges): s3.unmerge_cells(str(mr))
        for row in s3.iter_rows():
            for c in row: c.value = None
        s3._images = []
        thin = Side(style='thin', color='000000'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
        s3['B2'] = d['map_title']; s3['B2'].font = Font(bold=True, size=14)
        heads = ['입 력', 'Process Flow', '출 력', '책임부문', '성과지표']
        widths = [22, 34, 22, 12, 18]
        for j, (h, w) in enumerate(zip(heads, widths)):
            c = s3.cell(4, 2 + j, h); c.font = Font(bold=True); c.alignment = Alignment(horizontal='center', vertical='center'); c.border = box
            s3.column_dimensions[get_column_letter(2 + j)].width = w
        for i, f in enumerate(d['flow']):
            rr = 5 + i * 2
            for j, v in enumerate(f):
                c = s3.cell(rr, 2 + j, v); c.border = box; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            s3.row_dimensions[rr].height = 36
            if i < len(d['flow']) - 1:
                a = s3.cell(rr + 1, 3, '▼'); a.alignment = Alignment(horizontal='center')
    out = os.path.join(OUT, f"{d['code']}_{d['title']}.xlsx"); wb.save(out); return out

# ───────── 양식 ─────────
FORMS = [
  ('MI-0703-001', '클린룸 5S 점검표', '매일', ['점검 항목', '점검 기준', '결과(O/X/△)', '조치 내용', '확인'],
   [('정리', '보관 외 물품이 없을 것'), ('정돈', '지정 위치 보관, 미관상 양호'), ('청소', '바닥·작업대 청소 상태 및 주기 준수'), ('청결', '보관 Box·선반·자재 오염 없음'), ('습관화', '점검표 운용, 출입 기준 준수'), ('출입문·Air Shower', '작동 정상, 문 닫힘 유지')]),
  ('MI-0703-002', 'Particle 측정 관리대장', '매월', ['측정일', '측정 지점', '입자 크기(㎛)', 'Spec', '측정값', '판정', '측정자', '비고'], []),
  ('MI-0703-003', '클린복·보호구 관리 점검표', '매주', ['항목', '점검 기준', '결과(O/X)', '조치 내용', '확인'],
   [('클린복', '착용 상태, 오염·손상 없음, 세탁 주기 준수'), ('클린화', '전용 신발 착용, 외부 반출 금지'), ('모자·마스크', '머리카락 노출 없음'), ('장갑', '파손·오염 시 즉시 교체'), ('세탁·교체 기록', '세탁 일자 기록 유지')]),
  ('MI-0703-004', '검사실 온·습도 관리표', '매일', ['일자', '시간', '온도(℃)', '습도(%)', '판정 (20±2℃ / 65%↓)', '확인자', '비고'], []),
  ('MD-0701-004', 'Air Utility·필터 교체 관리대장', '매월', ['설비/유틸리티', 'Spec', '측정값', '필터/소모품', '교체 주기', '최근 교체일', '판정', '확인'],
   [('컴프레서 압축공기', ''), ('에어 드라이어', ''), ('라인 필터', ''), ('클린룸 FFU/HEPA 필터', '')]),
  ('MI-0804-002', 'Jig·Tool 정기 점검 Check Sheet', '매월', ['점검 항목', '점검 기준', '결과(OK/NG)', '조치 내용', '확인'],
   [('외관', '파손·마모·녹 없음'), ('치수/정도', '기준 치수 이내'), ('식별 표시', '번호·유효기간 라벨 부착'), ('보관 상태', '지정 위치 보관')]),
  ('MI-0807-004', '협력사 성적서 관리대장', '수시', ['입고일', '협력사', '품번/품명', 'LOT', '성적서 접수', '검토 결과', '확인자', '비고'], []),
  ('MD-0809-001', '협력사 개선 대책서', '수시', ['No', '지적사항', '원인', '개선 대책', '완료 예정일', '담당'], []),
  ('MD-0809-002', '협력사 개선 이행 점검 결과서', '수시', ['No', '지적사항', '개선 대책', '점검 결과', '판정(완료/미완료)', '점검일'], []),
  ('MI-0810-001', '고객 도면 배포·회수 대장', '수시', ['도면번호', 'Rev', '고객', '배포 부서', '배포일', '회수일', '구본 처리', '확인'], []),
  ('MI-0810-002', '정보보안·도면 관리 점검표', '매월', ['점검 항목', '결과(양호/미흡)', '조치 내용', '확인'],
   [('현장에 도면이 방치되어 있지 않은가',), ('도면이 성적서로 활용되고 있지 않은가',), ('구본 도면이 회수·폐기되었는가',), ('고객 도면 파일의 외부 반출·저장 매체 통제가 되는가',), ('출입자(방문객) 통제 기록이 있는가',)]),
  ('MI-0701-003', '외부 출처 문서 관리대장', '년', ['문서명', '발행처', '번호/버전', '발행·갱신일', '내용 확인', '보관 위치', '비고'],
   [('사업자등록증',), ('ISO 9001 인증서',), ('고객 품질 요구사항 (세메스 SSQ 등)',), ('관련 법규·KS 표준',)]),
  ('MI-0809-001', '샘플링 검사 기준표', '수시', ['로트 크기', '시료 문자', '시료 수', 'AQL', 'Ac (합격)', 'Re (불합격)'],
   [('2 ~ 8', 'A', '2'), ('9 ~ 15', 'B', '3'), ('16 ~ 25', 'C', '5'), ('26 ~ 50', 'D', '8'), ('51 ~ 90', 'E', '13'), ('91 ~ 150', 'F', '20'), ('151 ~ 280', 'G', '32'), ('281 ~ 500', 'H', '50'), ('501 ~ 1,200', 'J', '80'), ('1,201 ~ 3,200', 'K', '125')]),
]

def _style(c):
    return dict(font=copy.copy(c.font), fill=copy.copy(c.fill), border=copy.copy(c.border), alignment=copy.copy(c.alignment))

def _apply(cell, st):
    cell.font, cell.fill, cell.border, cell.alignment = st['font'], st['fill'], st['border'], st['alignment']

def _spans(n, first_wide):
    """32칸 격자를 n개 열로 나눔 (MST 원본 양식과 같은 격자)"""
    if n == 1: return [32]
    if first_wide:
        w0 = 10 if n <= 5 else 8
        rest = 32 - w0; base = rest // (n - 1); extra = rest - base * (n - 1)
        return [w0] + [base + (1 if i < extra else 0) for i in range(n - 1)]
    base = 32 // n; extra = 32 - base * n
    return [base + (1 if i < extra else 0) for i in range(n)]

def make_form(code, title, cycle, cols, items, logo, header=None):
    wb = openpyxl.load_workbook(TPL_FORM)
    ws = wb.worksheets[0]
    hdr_st = _style(ws['A9']); body_st = _style(ws['A11']); lab_st = _style(ws['A5']); val_st = _style(ws['D5'])
    ws.title = code
    ws['F2'] = title
    # 머리 항목 (원본 위치 A5/S5/A7/S7) — 양식번호·작성 주기를 기본으로
    header = header or [('양식번호', code), ('작성 주기', cycle), ('작성 부서', ''), ('작 성 일', '')]
    for (lab_cell, val_cell), hv in zip([('A5', 'D5'), ('S5', 'V5'), ('A7', 'D7'), ('S7', 'V7')], header + [('', '')] * 4):
        ws[lab_cell] = hv[0] or None; ws[val_cell] = hv[1] or None
    # 표 영역 다시 그리기 (9행부터)
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row >= 9: ws.unmerge_cells(str(mr))
    blank = Border()
    for row in ws.iter_rows(min_row=9, max_row=80, max_col=32):
        for c in row: c.value = None; c.border = blank
    spans = _spans(len(cols), bool(items))
    col = 1; starts = []
    for w in spans: starts.append(col); col += w
    def block(r1, r2, j, value, st):
        c1, c2 = starts[j], starts[j] + spans[j] - 1
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1): _apply(ws.cell(rr, cc), st)
        ws.cell(r1, c1, value)
        if c2 > c1 or r2 > r1: ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    for j, h in enumerate(cols): block(9, 10, j, h, hdr_st)
    nrows = max(len(items), 10)
    for i in range(nrows):
        r = 11 + i * 2
        it = items[i] if i < len(items) else ()
        for j in range(len(cols)):
            block(r, r + 1, j, it[j] if j < len(it) else None, body_st)
    foot = 11 + nrows * 2 + 1
    ws.cell(foot, 1, f'{code}  ·  주식회사 엠에스티  ·  보존기간 3년').font = Font(size=8)
    out = os.path.join(OUT, f'{code}_{title}.xlsx'); wb.save(out); return out

if __name__ == '__main__':
    logo = logo_bytes(TPL_DOC)
    for d in DOCS: print(make_doc(d, logo))
    for f in FORMS: print(make_form(*f, logo))
