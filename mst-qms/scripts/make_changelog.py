"""원본 대비 정정 내역 (고객 제출용) — 원본_정정본/_logs/*.json 을 하나의 엑셀로
사용: python3 make_changelog.py <원본_정정본 폴더> <출력 xlsx>"""
import glob, json, os, re, sys, unicodedata
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SRC, OUT = sys.argv[1], sys.argv[2]
thin = Side(style='thin', color='999999'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='DCE6F2'); F = '맑은 고딕'
files = {}
for p in glob.glob(os.path.join(SRC, '**', '*.*'), recursive=True):
    if '/_logs/' in p or not re.search(r'\.(xlsx|pptx)$', p): continue
    files[unicodedata.normalize('NFC', os.path.splitext(os.path.basename(p))[0])] = os.path.relpath(os.path.dirname(p), SRC)
rows, seen = [], set()
def add(fname, where, before, after):
    b, a = (before or '').strip(), (after or '').strip()
    if b == a: return
    key = (fname, b, a)
    if key in seen: return
    seen.add(key)
    rows.append([files.get(fname, ''), fname, where, b[:300], a[:300] if a else '(삭제)'])
for p in sorted(glob.glob(os.path.join(SRC, '_logs', '*.json'))):
    name = unicodedata.normalize('NFC', os.path.splitext(os.path.basename(p))[0])
    try: data = json.load(open(p, encoding='utf-8'))
    except Exception: continue
    if name.startswith('결정내역'):
        items = data if isinstance(data, list) else data.get('decisions', [])
        for x in items:
            fn = unicodedata.normalize('NFC', os.path.splitext(os.path.basename(x.get('file', '')))[0])
            add(fn, x.get('where', '') or '', x.get('before', ''), x.get('after', ''))
        continue
    for c in data.get('changes', []):
        part = c.get('part', ''); m = re.search(r'(sheet|slide)(\d+)', part)
        add(name, (('시트 ' if m and m.group(1) == 'sheet' else '슬라이드 ') + m.group(2)) if m else '', c.get('before', ''), c.get('after', ''))
    for c in data.get('cells', []):
        if 'error' not in c: add(name, f"{c.get('sheet', '')} {c.get('cell', '')}", c.get('before', ''), c.get('after', ''))
wb = Workbook(); ws = wb.active; ws.title = '원본 대비 정정 내역'
ws['A1'] = '주식회사 엠에스티 품질경영시스템 문서 — 원본 대비 정정 내역'; ws['A1'].font = Font(name=F, size=14, bold=True)
ws['A2'] = f'정정 {len(rows)}건 · 대상 파일 {len(set(r[1] for r in rows))}종 · 오타·맞춤법, 타사 명칭·번호, 문서번호 체계, 업종 표현 정정'; ws['A2'].font = Font(name=F, size=10, color='555555')
heads = ['폴더', '문서(파일)', '위치', '정정 전', '정정 후']
for j, h in enumerate(heads, 1):
    c = ws.cell(4, j, h); c.font = Font(name=F, bold=True); c.fill = HEAD; c.border = BOX; c.alignment = Alignment(horizontal='center')
rows.sort(key=lambda r: (r[0], r[1]))
for i, r in enumerate(rows, 5):
    for j, v in enumerate(r, 1):
        c = ws.cell(i, j, v); c.font = Font(name=F, size=9); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical='top')
for col, w in zip('ABCDE', [26, 36, 12, 60, 60]): ws.column_dimensions[col].width = w
ws.freeze_panes = 'A5'
wb.save(OUT); print('rows', len(rows))
if len(sys.argv) > 3:   # 프로그램 화면용 데이터
    open(sys.argv[3], 'w', encoding='utf-8').write('/* 자동 생성: scripts/make_changelog.py */\nwindow.SEED = window.SEED || {};\nwindow.SEED.revisionLog = ' + json.dumps([{'kind': r[0].split('/')[0] if r[0] else '', 'file': r[1], 'where': r[1] + (' · ' + r[2] if r[2] else ''), 'before': r[3], 'after': r[4]} for r in rows], ensure_ascii=False) + ';\n')
