# -*- coding: utf-8 -*-
"""원본_정정본의 xlsx/pptx → app/preview/*.html (앱 안에서 바로 열람)
   + app/data/previews.js (원본 경로 → 미리보기 경로, 시트 목록)
   + app/data/turtle.js (MP 프로세스 '터틀분석' 시트 도형 → 터틀맵 데이터)
   셀 서식(병합·열폭·행높이·테두리·채우기·글꼴·정렬), 그림, 도형(그룹 포함), 표(pptx) 렌더."""
import os, re, io, json, hashlib, zipfile, unicodedata, datetime, html, posixpath, glob
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, '원본_정정본')
OUT = os.path.join(ROOT, 'app', 'preview')
IMG = os.path.join(OUT, 'img')
EMU = 9525.0

def nfc(s): return unicodedata.normalize('NFC', s)
def esc(s): return html.escape(str(s), quote=True)
def loc(t): return t.split('}')[-1]
def kids(e, name): return [c for c in e if loc(c.tag) == name]
def kid(e, name):
    if e is None: return None
    for c in e:
        if loc(c.tag) == name: return c
    return None
def path_(e, *names):
    for n in names:
        e = kid(e, n)
        if e is None: return None
    return e
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

# ───────────── 색상 ─────────────
INDEXED = ['000000','FFFFFF','FF0000','00FF00','0000FF','FFFF00','FF00FF','00FFFF','000000','FFFFFF','FF0000','00FF00','0000FF','FFFF00','FF00FF','00FFFF','800000','008000','000080','808000','800080','008080','C0C0C0','808080','9999FF','993366','FFFFCC','CCFFFF','660066','FF8080','0066CC','CCCCFF','000080','FF00FF','FFFF00','00FFFF','800080','800000','008080','0000FF','00CCFF','CCFFFF','CCFFCC','FFFF99','99CCFF','FF99CC','CC99FF','FFCC99','3366FF','33CCCC','99CC00','FFCC00','FF9900','FF6600','666699','969696','003366','339966','003300','333300','993300','993366','333399','333333']

def tint(hexc, t):
    if not t: return hexc
    r, g, b = int(hexc[0:2], 16), int(hexc[2:4], 16), int(hexc[4:6], 16)
    def f(c): return int(c * (1 + t)) if t < 0 else int(c + (255 - c) * t)
    return '%02X%02X%02X' % (f(r), f(g), f(b))

def lummod(hexc, mod, off):
    import colorsys
    r, g, b = [int(hexc[i:i+2], 16) / 255 for i in (0, 2, 4)]
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = min(1, max(0, l * mod + off))
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return '%02X%02X%02X' % (int(r * 255), int(g * 255), int(b * 255))

def theme_colors(z, name):
    try: x = ET.fromstring(z.read(name))
    except KeyError: return {}
    cs = None
    for e in x.iter():
        if loc(e.tag) == 'clrScheme': cs = e; break
    out = {}
    if cs is None: return out
    for c in cs:
        v = None
        for s in c:
            if loc(s.tag) == 'srgbClr': v = s.get('val')
            elif loc(s.tag) == 'sysClr': v = s.get('lastClr') or ('000000' if s.get('val') == 'windowText' else 'FFFFFF')
        out[loc(c.tag)] = v
    return out

class Theme:
    def __init__(self, tc):
        self.tc = tc
        self.idx = [tc.get(k, 'FFFFFF') for k in ('lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6', 'hlink', 'folHlink')]
    def xl(self, col):
        """openpyxl Color → '#RRGGBB' | None"""
        if col is None: return None
        try:
            if col.type == 'rgb':
                v = col.rgb
                if not isinstance(v, str) or len(v) < 6: return None
                if len(v) == 8 and v[:2] == '00' and v[2:] == '000000' and False: return None
                return '#' + v[-6:]
            if col.type == 'theme':
                base = self.idx[col.theme] if col.theme < len(self.idx) else '000000'
                return '#' + tint(base, col.tint or 0)
            if col.type == 'indexed':
                if col.indexed in (64,): return None
                if col.indexed < len(INDEXED): return '#' + INDEXED[col.indexed]
        except Exception: return None
        return None
    def dml(self, e):
        """DrawingML 색 요소(srgbClr/schemeClr/prstClr/sysClr) → '#RRGGBB'"""
        if e is None: return None
        t = loc(e.tag); v = None
        if t == 'srgbClr': v = e.get('val')
        elif t == 'sysClr': v = e.get('lastClr') or '000000'
        elif t == 'prstClr': v = {'black': '000000', 'white': 'FFFFFF', 'red': 'FF0000', 'blue': '0000FF', 'green': '00FF00', 'yellow': 'FFFF00'}.get(e.get('val'), '000000')
        elif t == 'schemeClr':
            k = e.get('val'); k = {'tx1': 'dk1', 'bg1': 'lt1', 'tx2': 'dk2', 'bg2': 'lt2', 'phClr': 'accent1'}.get(k, k)
            v = self.tc.get(k) or '000000'
        if not v: return None
        mod, off = 1.0, 0.0; alpha = None
        for m in e:
            n = loc(m.tag)
            if n == 'lumMod': mod = int(m.get('val')) / 100000
            elif n == 'lumOff': off = int(m.get('val')) / 100000
            elif n == 'tint': v = tint(v, 1 - int(m.get('val')) / 100000)
            elif n == 'shade': v = tint(v, -(1 - int(m.get('val')) / 100000))
            elif n == 'alpha': alpha = int(m.get('val')) / 100000
        if mod != 1.0 or off: v = lummod(v, mod, off)
        if alpha is not None and alpha < 0.99:
            return 'rgba(%d,%d,%d,%.2f)' % (int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16), alpha)
        return '#' + v
    def fill_of(self, el):
        """spPr/tcPr 등에서 채우기"""
        if el is None: return None
        for c in el:
            n = loc(c.tag)
            if n == 'noFill': return 'none'
            if n == 'solidFill': return self.dml(c[0]) if len(c) else None
            if n == 'gradFill':
                gs = path_(c, 'gsLst')
                if gs is not None and len(gs): return self.dml(gs[0][0]) if len(gs[0]) else None
        return None

# ───────────── 이미지 저장 ─────────────
def save_img(data, ext):
    h = hashlib.sha1(data).hexdigest()[:16]
    ext = ext.lower().lstrip('.')
    if ext in ('jpeg',): ext = 'jpg'
    if ext in ('emf', 'wmf', 'tif', 'tiff'):
        try:
            from PIL import Image
            im = Image.open(io.BytesIO(data)); b = io.BytesIO(); im.save(b, 'PNG'); data = b.getvalue(); ext = 'png'
        except Exception: return None
    fn = h + '.' + ext
    p = os.path.join(IMG, fn)
    if not os.path.exists(p):
        with open(p, 'wb') as f: f.write(data)
    return 'img/' + fn

def rels_of(z, part):
    d, b = posixpath.split(part)
    rp = posixpath.join(d, '_rels', b + '.rels')
    out = {}
    try: x = ET.fromstring(z.read(rp))
    except KeyError: return out
    for r in x:
        t = r.get('Target')
        if r.get('TargetMode') == 'External': continue
        out[r.get('Id')] = posixpath.normpath(posixpath.join(d, t)) if not t.startswith('/') else t.lstrip('/')
    return out

# ───────────── 도형 렌더 (DrawingML 공통) ─────────────
GEOM = {
    'hexagon': '25,0 75,0 100,50 75,100 25,100 0,50',
    'octagon': '29,0 71,0 100,29 100,71 71,100 29,100 0,71 0,29',
    'diamond': '50,0 100,50 50,100 0,50', 'flowChartDecision': '50,0 100,50 50,100 0,50',
    'triangle': '50,0 100,100 0,100', 'rtTriangle': '0,0 100,100 0,100',
    'parallelogram': '25,0 100,0 75,100 0,100', 'flowChartInputOutput': '20,0 100,0 80,100 0,100',
    'trapezoid': '25,0 75,0 100,100 0,100',
    'pentagon': '50,0 100,38 81,100 19,100 0,38',
    'homePlate': '0,0 80,0 100,50 80,100 0,100', 'chevron': '0,0 80,0 100,50 80,100 0,100 20,50',
    'rightArrow': '0,25 70,25 70,0 100,50 70,100 70,75 0,75', 'leftArrow': '100,25 30,25 30,0 0,50 30,100 30,75 100,75',
    'downArrow': '25,0 75,0 75,70 100,70 50,100 0,70 25,70', 'upArrow': '25,100 75,100 75,30 100,30 50,0 0,30 25,30',
    'leftRightArrow': '0,50 20,0 20,25 80,25 80,0 100,50 80,100 80,75 20,75 20,100',
    'upDownArrow': '50,0 100,20 75,20 75,80 100,80 50,100 0,80 25,80 25,20 0,20',
    'flowChartPredefinedProcess': None, 'flowChartManualInput': '0,20 100,0 100,100 0,100',
    'flowChartOffpageConnector': '0,0 100,0 100,80 50,100 0,80', 'snip1Rect': '0,0 85,0 100,15 100,100 0,100',
    'notchedRightArrow': '0,25 70,25 70,0 100,50 70,100 70,75 0,75 15,50',
    'stripedRightArrow': '0,25 70,25 70,0 100,50 70,100 70,75 0,75',
}
DASH = {'dash': '6,4', 'sysDash': '4,3', 'dot': '1,3', 'sysDot': '1,2', 'lgDash': '10,4', 'dashDot': '6,3,1,3'}

class Ctx:
    def __init__(self, z, part, theme, kind):
        self.z, self.part, self.theme, self.kind = z, part, theme, kind
        self.rels = rels_of(z, part)
        self.default_sz = 11.0 if kind == 'xlsx' else 18.0

def text_html(ctx, tx, default_color=None, lst_style=None, default_sz=None, default_algn=None, default_b=False):
    """txBody → HTML (문단·런)"""
    if tx is None: return ''
    out = []
    for p in kids(tx, 'p'):
        ppr = kid(p, 'pPr')
        algn = (ppr.get('algn') if ppr is not None else None) or default_algn or 'l'
        al = {'l': 'left', 'ctr': 'center', 'r': 'right', 'just': 'justify', 'dist': 'justify'}.get(algn, 'left')
        bullet = ''
        if ppr is not None:
            bc = kid(ppr, 'buChar')
            if bc is not None: bullet = esc(bc.get('char', '•')) + ' '
            if kid(ppr, 'buAutoNum') is not None: bullet = ''
        lvl_indent = int(ppr.get('lvl', '0')) * 14 if ppr is not None else 0
        end = kid(p, 'endParaRPr')
        runs = []
        psz = None
        for r in p:
            n = loc(r.tag)
            if n == 'br': runs.append('<br>'); continue
            if n not in ('r', 'fld'): continue
            t = kid(r, 't'); t = t.text if t is not None and t.text else ''
            if n == 'fld' and (r.get('type') or '').startswith('slidenum') and getattr(ctx, 'slide_no', None): t = str(ctx.slide_no)
            rpr = kid(r, 'rPr')
            st = []
            sz = rpr.get('sz') if rpr is not None else None
            if sz: st.append('font-size:%.1fpt' % (int(sz) / 100)); psz = psz or int(sz) / 100
            b = rpr.get('b') if rpr is not None else None
            if b == '1' or (b is None and default_b): st.append('font-weight:700')
            if rpr is not None and rpr.get('i') == '1': st.append('font-style:italic')
            if rpr is not None and rpr.get('u') not in (None, 'none'): st.append('text-decoration:underline')
            col = None
            if rpr is not None:
                sf = kid(rpr, 'solidFill')
                if sf is not None and len(sf): col = ctx.theme.dml(sf[0])
            col = col or default_color
            if col: st.append('color:' + col)
            runs.append('<span style="%s">%s</span>' % (';'.join(st), esc(t).replace('  ', ' &nbsp;')))
        body = ''.join(runs)
        if not body.strip():
            esz = end.get('sz') if end is not None else None
            out.append('<p style="font-size:%.1fpt">&nbsp;</p>' % ((int(esz) / 100) if esz else (default_sz or ctx.default_sz)))
            continue
        out.append('<p style="text-align:%s;padding-left:%dpx">%s%s</p>' % (al, lvl_indent, bullet, body))
    return ''.join(out)

def style_ref_color(ctx, sp, name):
    st = kid(sp, 'style')
    if st is None: return None
    r = kid(st, name)
    if r is None or r.get('idx') == '0' or not len(r): return None
    return ctx.theme.dml(r[0])

def render_sp(ctx, sp, box, is_cxn=False, ph_style=None):
    """box=(x,y,w,h) px (부모 좌표계)"""
    x, y, w, h = box
    sppr = kid(sp, 'spPr')
    xfrm = kid(sppr, 'xfrm') if sppr is not None else None
    rot = int(xfrm.get('rot', '0')) / 60000 if xfrm is not None else 0
    fh = xfrm is not None and xfrm.get('flipH') == '1'
    fv = xfrm is not None and xfrm.get('flipV') == '1'
    geom = path_(sppr, 'prstGeom') if sppr is not None else None
    prst = geom.get('prst') if geom is not None else ('rect' if kid(sppr, 'custGeom') is None else 'cust')
    fill = ctx.theme.fill_of(sppr)
    if fill is None: fill = style_ref_color(ctx, sp, 'fillRef')
    if fill == 'none': fill = None
    ln = kid(sppr, 'ln') if sppr is not None else None
    lc, lw, dash = None, 1.0, None
    if ln is not None:
        if kid(ln, 'noFill') is not None: lc = 'none'
        else:
            sf = kid(ln, 'solidFill')
            if sf is not None and len(sf): lc = ctx.theme.dml(sf[0])
        if ln.get('w'): lw = max(0.75, int(ln.get('w')) / EMU)
        pd = kid(ln, 'prstDash')
        if pd is not None: dash = DASH.get(pd.get('val'))
    if lc is None: lc = style_ref_color(ctx, sp, 'lnRef')
    if lc == 'none': lc = None
    tf = 'transform:' + ' '.join(t for t in ['rotate(%.1fdeg)' % rot if rot else '', 'scaleX(-1)' if fh and (is_cxn or prst != 'rect') else '', 'scaleY(-1)' if fv and (is_cxn or prst != 'rect') else ''] if t) + ';' if (rot or fh or fv) else ''
    base = 'left:%.1fpx;top:%.1fpx;width:%.1fpx;height:%.1fpx;%s' % (x, y, max(w, 1), max(h, 1), tf)
    if is_cxn or prst in ('line', 'straightConnector1', 'bentConnector2', 'bentConnector3', 'curvedConnector3'):
        c = lc or '#000'
        head = path_(sppr, 'ln', 'headEnd') if sppr is not None else None
        tail = path_(sppr, 'ln', 'tailEnd') if sppr is not None else None
        mk = ''
        mid = 'm' + hashlib.md5(('%s%s%s' % (x, y, c)).encode()).hexdigest()[:6]
        if (tail is not None and tail.get('type') not in (None, 'none')) or (head is not None and head.get('type') not in (None, 'none')):
            mk = '<defs><marker id="%s" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="6" markerHeight="6" orient="auto-start-reverse"><path d="M0,0L10,5L0,10z" fill="%s"/></marker></defs>' % (mid, c)
        ms = ' marker-end="url(#%s)"' % mid if tail is not None and tail.get('type') not in (None, 'none') else ''
        mh = ' marker-start="url(#%s)"' % mid if head is not None and head.get('type') not in (None, 'none') else ''
        if prst == 'bentConnector3':
            d = 'M0,0 L%.1f,0 L%.1f,%.1f L%.1f,%.1f' % (w / 2, w / 2, h, w, h)
        else:
            d = 'M0,0 L%.1f,%.1f' % (max(w, 0.5), max(h, 0.5))
        return '<svg class="shp" style="%soverflow:visible" width="%.1f" height="%.1f">%s<path d="%s" fill="none" stroke="%s" stroke-width="%.1f"%s%s%s/></svg>' % (base, max(w, 1), max(h, 1), mk, d, c, lw, ' stroke-dasharray="%s"' % dash if dash else '', ms, mh)
    # 본체
    parts = []
    st = base
    if prst in ('rect', 'cust', 'flowChartProcess', 'flowChartAlternateProcess', 'roundRect', 'ellipse', 'flowChartConnector', 'flowChartTerminator', 'snipRoundRect', 'round2SameRect', 'plaque', 'can', 'flowChartDocument', 'flowChartPredefinedProcess', 'bracketPair', 'foldedCorner', 'frame', 'bevel', 'cube') or prst not in GEOM:
        if fill: st += 'background:%s;' % fill
        if lc: st += 'border:%.1fpx %s %s;' % (lw, 'dashed' if dash else 'solid', lc)
        if prst in ('roundRect', 'flowChartAlternateProcess', 'snipRoundRect', 'round2SameRect', 'plaque'): st += 'border-radius:%dpx;' % int(min(w, h) * 0.16)
        if prst in ('ellipse', 'flowChartConnector'): st += 'border-radius:50%;'
        if prst == 'flowChartTerminator': st += 'border-radius:%dpx;' % int(h / 2)
    else:
        pts = GEOM[prst]
        parts.append('<svg class="geo" viewBox="0 0 100 100" preserveAspectRatio="none"><polygon points="%s" fill="%s" stroke="%s" stroke-width="%.1f" vector-effect="non-scaling-stroke"/></svg>' % (pts, fill or 'none', lc or 'none', lw))
    tx = kid(sp, 'txBody')
    if tx is not None:
        bp = kid(tx, 'bodyPr')
        anchor = (bp.get('anchor') if bp is not None else None) or (ph_style or {}).get('anchor') or 't'
        jc = {'t': 'flex-start', 'ctr': 'center', 'b': 'flex-end'}.get(anchor, 'flex-start')
        vert = bp is not None and bp.get('vert') in ('vert', 'eaVert', 'wordArtVert', 'vert270')
        li = int(bp.get('lIns', '91440')) / EMU if bp is not None else 9.6
        ti = int(bp.get('tIns', '45720')) / EMU if bp is not None else 4.8
        tcol = style_ref_color(ctx, sp, 'fontRef')
        inner = text_html(ctx, tx, tcol, default_sz=(ph_style or {}).get('sz'), default_algn=(ph_style or {}).get('algn'), default_b=(ph_style or {}).get('b'))
        wrap = 'nowrap' if (bp is not None and bp.get('wrap') == 'none') else 'normal'
        parts.append('<div class="tx" style="justify-content:%s;padding:%.1fpx %.1fpx;white-space:%s;%s%s">%s</div>' % (jc, ti, li, wrap, 'writing-mode:vertical-rl;' if vert else '', ('font-size:%.1fpt;' % ph_style['sz']) if ph_style and ph_style.get('sz') else '', inner))
    return '<div class="shp" style="%s">%s</div>' % (st, ''.join(parts))

def render_pic(ctx, pic, box):
    x, y, w, h = box
    blip = None
    for e in pic.iter():
        if loc(e.tag) == 'blip': blip = e; break
    if blip is None: return ''
    rid = blip.get('{%s}embed' % R_NS)
    tgt = ctx.rels.get(rid)
    if not tgt: return ''
    try: data = ctx.z.read(tgt)
    except KeyError: return ''
    src = save_img(data, posixpath.splitext(tgt)[1])
    if not src: return ''
    sppr = kid(pic, 'spPr'); xfrm = kid(sppr, 'xfrm') if sppr is not None else None
    rot = int(xfrm.get('rot', '0')) / 60000 if xfrm is not None else 0
    return '<img class="shp" src="%s" style="left:%.1fpx;top:%.1fpx;width:%.1fpx;height:%.1fpx;%s" alt="">' % (src, x, y, w, h, 'transform:rotate(%.1fdeg);' % rot if rot else '')

def render_tbl(ctx, gf, box):
    x, y, w, h = box
    tbl = None
    for e in gf.iter():
        if loc(e.tag) == 'tbl': tbl = e; break
    if tbl is None: return '<div class="shp chart" style="left:%.1fpx;top:%.1fpx;width:%.1fpx;height:%.1fpx">[차트/개체]</div>' % (x, y, w, h)
    grid = [int(g.get('w')) / EMU for g in kids(kid(tbl, 'tblGrid'), 'gridCol')]
    tot = sum(grid) or 1
    sc = w / tot if w else 1
    rows = []
    for tr in kids(tbl, 'tr'):
        cells = []
        for tc in kids(tr, 'tc'):
            if tc.get('hMerge') == '1' or tc.get('vMerge') == '1': continue
            span = ''
            if tc.get('gridSpan'): span += ' colspan="%s"' % tc.get('gridSpan')
            if tc.get('rowSpan'): span += ' rowspan="%s"' % tc.get('rowSpan')
            tcpr = kid(tc, 'tcPr')
            st = []
            f = ctx.theme.fill_of(tcpr)
            if f and f != 'none': st.append('background:' + f)
            if tcpr is not None:
                for side, nm in (('left', 'lnL'), ('right', 'lnR'), ('top', 'lnT'), ('bottom', 'lnB')):
                    l = kid(tcpr, nm)
                    if l is not None:
                        if kid(l, 'noFill') is not None: st.append('border-%s:none' % side)
                        else:
                            sf = kid(l, 'solidFill')
                            c = ctx.theme.dml(sf[0]) if sf is not None and len(sf) else '#000'
                            st.append('border-%s:%.1fpx solid %s' % (side, max(0.75, int(l.get('w', '12700')) / EMU), c))
                anc = tcpr.get('anchor')
                st.append('vertical-align:' + {'ctr': 'middle', 'b': 'bottom'}.get(anc, 'top'))
            cells.append('<td%s style="%s">%s</td>' % (span, ';'.join(st), text_html(ctx, kid(tc, 'txBody'), default_sz=10.0)))
        rows.append('<tr style="height:%.1fpx">%s</tr>' % (int(tr.get('h', '0')) / EMU * sc, ''.join(cells)))
    cols = ''.join('<col style="width:%.1fpx">' % (g * sc) for g in grid)
    return '<table class="shp ptbl" style="left:%.1fpx;top:%.1fpx;width:%.1fpx"><colgroup>%s</colgroup>%s</table>' % (x, y, w, cols, ''.join(rows))

def xfrm_of(el):
    """sp/pic/grpSp/cxnSp/graphicFrame 의 xfrm (off, ext, chOff, chExt)"""
    n = loc(el.tag)
    pr = kid(el, 'grpSpPr') if n == 'grpSp' else (kid(el, 'xfrm') if n == 'graphicFrame' else kid(el, 'spPr'))
    xf = pr if (n == 'graphicFrame') else (kid(pr, 'xfrm') if pr is not None else None)
    if xf is None: return None
    o, e = kid(xf, 'off'), kid(xf, 'ext')
    if o is None or e is None: return None
    r = {'x': int(o.get('x')), 'y': int(o.get('y')), 'cx': int(e.get('cx')), 'cy': int(e.get('cy'))}
    co, ce = kid(xf, 'chOff'), kid(xf, 'chExt')
    if co is not None and ce is not None:
        r.update(chx=int(co.get('x')), chy=int(co.get('y')), chcx=int(ce.get('cx')) or 1, chcy=int(ce.get('cy')) or 1)
    return r

def render_el(ctx, el, T, ph_lookup=None):
    """T: (x_emu, y_emu) → (px, py) 스케일 포함 함수 (T, sx, sy)"""
    fn, sx, sy = T
    n = loc(el.tag)
    xf = xfrm_of(el)
    ph_style = None
    if xf is None and ph_lookup is not None and n == 'sp':
        xf, ph_style = ph_lookup(el)
    if xf is None: return ''
    px, py = fn(xf['x'], xf['y'])
    box = (px, py, xf['cx'] * sx, xf['cy'] * sy)
    return render_el_box(ctx, el, box, xf, T, ph_style)

def render_el_box(ctx, el, box, xf, T=None, ph_style=None):
    n = loc(el.tag)
    if n == 'sp': return render_sp(ctx, el, box, ph_style=ph_style)
    if n == 'cxnSp': return render_sp(ctx, el, box, is_cxn=True)
    if n == 'pic': return render_pic(ctx, el, box)
    if n == 'graphicFrame': return render_tbl(ctx, el, box)
    if n == 'grpSp':
        if xf is None or 'chx' not in xf: return ''
        gx, gy, gw, gh = box
        sx = gw / xf['chcx']; sy = gh / xf['chcy']
        def fn(x, y, xf=xf, gx=gx, gy=gy, sx=sx, sy=sy): return (gx + (x - xf['chx']) * sx, gy + (y - xf['chy']) * sy)
        inner = ''.join(render_el(ctx, c, (fn, sx, sy)) for c in el if loc(c.tag) in ('sp', 'cxnSp', 'pic', 'grpSp', 'graphicFrame'))
        return inner
    if n == 'AlternateContent':
        for c in el.iter():
            if loc(c.tag) == 'Fallback':
                return ''.join(render_el_box(ctx, cc, box, xfrm_of(cc) or xf, T) for cc in c)
    return ''

# ───────────── XLSX ─────────────
def num_fmt(v, fmt):
    if isinstance(v, (datetime.datetime, datetime.date)):
        if isinstance(v, datetime.datetime) and (v.hour or v.minute) and ('h' in fmt.lower()):
            return v.strftime('%Y-%m-%d %H:%M')
        f = fmt.lower()
        if 'mm-dd' in f and 'yy' not in f: return v.strftime('%m-%d')
        if 'yy' in f and '"년"' in fmt: return '%d년 %d월 %d일' % (v.year, v.month, v.day)
        return v.strftime('%Y-%m-%d')
    if isinstance(v, datetime.time): return v.strftime('%H:%M')
    if isinstance(v, bool): return 'TRUE' if v else 'FALSE'
    if isinstance(v, (int, float)):
        f = (fmt or 'General').split(';')[0]
        pct = '%' in f
        x = v * 100 if pct else v
        m = re.search(r'0\.(0+)', f)
        dec = len(m.group(1)) if m else (0 if re.search(r'[0#]', f) and f != 'General' else None)
        if dec is None:
            if isinstance(x, float):
                s = ('%.10g' % x)
                if 'e' in s: s = '%.6g' % x
            else: s = str(x)
        else:
            s = ('{:,.%df}' % dec).format(x) if ',' in f else ('{:.%df}' % dec).format(x)
        return s + ('%' if pct else '')
    return str(v)

BSTY = {'thin': (1, 'solid'), 'medium': (2, 'solid'), 'thick': (3, 'solid'), 'dashed': (1, 'dashed'), 'dotted': (1, 'dotted'), 'double': (3, 'double'), 'hair': (1, 'dotted'), 'mediumDashed': (2, 'dashed'), 'dashDot': (1, 'dashed'), 'mediumDashDot': (2, 'dashed'), 'dashDotDot': (1, 'dashed'), 'mediumDashDotDot': (2, 'dashed'), 'slantDashDot': (2, 'dashed')}

def bcss(side, b, theme):
    if b is None or not b.style: return ''
    w, s = BSTY.get(b.style, (1, 'solid'))
    c = theme.xl(b.color) or '#000'
    return 'border-%s:%dpx %s %s;' % (side, w, s, c)

def sheet_parts(z):
    """시트 이름 → (sheet part, drawing part|None)"""
    wb = ET.fromstring(z.read('xl/workbook.xml'))
    wr = rels_of(z, 'xl/workbook.xml')
    out = {}
    for s in wb.iter():
        if loc(s.tag) == 'sheet':
            part = wr.get(s.get('{%s}id' % R_NS))
            if not part: continue
            dr = None
            try:
                sx = ET.fromstring(z.read(part))
                d = None
                for e in sx:
                    if loc(e.tag) == 'drawing': d = e.get('{%s}id' % R_NS)
                if d: dr = rels_of(z, part).get(d)
            except KeyError: pass
            out[s.get('name')] = (part, dr)
    return out

def anchors(z, dpart, theme):
    """drawing → [(kind, from, to|ext, element)]"""
    x = ET.fromstring(z.read(dpart))
    res = []
    for a in x:
        n = loc(a.tag)
        fr = kid(a, 'from'); to = kid(a, 'to')
        def pt(e): return (int(kid(e, 'col').text), int(kid(e, 'colOff').text) / EMU, int(kid(e, 'row').text), int(kid(e, 'rowOff').text) / EMU)
        el = None
        for c in a:
            if loc(c.tag) in ('sp', 'grpSp', 'pic', 'cxnSp', 'graphicFrame', 'AlternateContent'): el = c
        if el is None: continue
        if n == 'twoCellAnchor' and fr is not None and to is not None: res.append(('two', pt(fr), pt(to), el))
        elif n == 'oneCellAnchor' and fr is not None:
            e = kid(a, 'ext'); res.append(('one', pt(fr), (int(e.get('cx')) / EMU, int(e.get('cy')) / EMU), el))
        elif n == 'absoluteAnchor':
            p = kid(a, 'pos'); e = kid(a, 'ext'); res.append(('abs', (int(p.get('x')) / EMU, int(p.get('y')) / EMU), (int(e.get('cx')) / EMU, int(e.get('cy')) / EMU), el))
    return res

def col_px(w):
    if w is None: return 64
    return int(((256 * w + int(128 / 7)) / 256) * 7)

TURTLE_CUR = {}
def turtle_html(t, w, h):
    """원본 터틀분석 도형 → 정돈된 터틀 다이어그램 (내용은 원본 그대로)"""
    def box(k, cls):
        items = ''.join('<li>%s</li>' % esc(v) for v in t.get(k, []))
        return '<div class="tt %s"><b>%s</b><ul>%s</ul></div>' % (cls, esc(t.get(k + 'Label', k)), items)
    return ('<div class="turtle" style="width:%.0fpx;height:%.0fpx">' % (w, h) + box('what', 'a') + box('who', 'b') + box('input', 'c')
            + '<div class="tt core"><svg viewBox="0 0 100 86" preserveAspectRatio="none"><polygon points="25,0 75,0 100,43 75,86 25,86 0,43" fill="#CCFFCC" stroke="#333" stroke-width="1" vector-effect="non-scaling-stroke"/></svg><span>%s<br>PROCESS</span></div>' % esc(t.get('name', ''))
            + box('output', 'd') + box('how', 'e') + box('measure', 'f') + '</div>')

def render_xlsx(path):
    z = zipfile.ZipFile(path)
    TURTLE_CUR['t'] = turtle_of(path) if re.match(r'^MP-\d{4}', nfc(os.path.basename(path))) else None
    theme = Theme(theme_colors(z, 'xl/theme/theme1.xml'))
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = sheet_parts(z)
    sheets = []
    for ws in wb.worksheets:
        if ws.sheet_state != 'visible': continue
        part, dpart = parts.get(ws.title, (None, None))
        ctx = Ctx(z, dpart or 'xl/drawings/x.xml', theme, 'xlsx')
        anc = anchors(z, dpart, theme) if dpart else []
        # 사용 범위
        maxr = maxc = 0
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and str(c.value).strip() != '':
                    maxr = max(maxr, c.row); maxc = max(maxc, c.column)
        for mr in ws.merged_cells.ranges:
            tl = ws.cell(mr.min_row, mr.min_col)
            if tl.value is not None or mr.min_row <= maxr:
                if tl.value is not None or (tl.has_style and tl.border and tl.border.left and tl.border.left.style):
                    maxr = max(maxr, mr.max_row); maxc = max(maxc, mr.max_col)
        # 테두리만 있는 빈 표(기입란)도 포함 — 값 있는 범위의 열 안에서 아래로 이어지는 테두리 행
        if maxc:
            r = maxr + 1
            while r <= min(ws.max_row, maxr + 400):
                has = False
                for cc in range(1, maxc + 1):
                    b = ws.cell(r, cc).border
                    if b is not None and ((b.left and b.left.style) or (b.bottom and b.bottom.style) or (b.top and b.top.style)): has = True; break
                if not has: break
                maxr = r; r += 1
        for k, fr, to, el in anc:
            if k == 'two': maxr = max(maxr, to[2] + 1); maxc = max(maxc, to[0] + 1)
            elif k == 'one': maxr = max(maxr, fr[2] + 2); maxc = max(maxc, fr[0] + 2)
        pa = None
        try:
            if ws.print_area:
                m = re.search(r'\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)', str(ws.print_area))
                if m:
                    from openpyxl.utils import column_index_from_string as ci
                    pa = (ci(m.group(1)), int(m.group(2)), ci(m.group(3)), int(m.group(4)))
        except Exception: pa = None
        if pa:
            maxc = max(min(maxc, pa[2]) if maxc else pa[2], pa[2]) if pa[2] <= 80 else maxc
            maxr = max(maxr, pa[3]) if pa[3] <= 600 else maxr
        maxr = min(maxr, 800); maxc = min(maxc, 80)
        if not maxr or not maxc:
            if not anc: continue
        dw = ws.sheet_format.defaultColWidth or (ws.sheet_format.baseColWidth or 8) + 0.71
        dh = ws.sheet_format.defaultRowHeight or 15
        colw = []
        for c in range(1, max(maxc, 1) + 1):
            cd = ws.column_dimensions.get(get_column_letter(c))
            hidden = False; w = None
            # column_dimensions 는 min~max 묶음일 수 있음
            for key, d in ws.column_dimensions.items():
                if d.min and d.max and d.min <= c <= d.max:
                    if d.hidden: hidden = True
                    if d.width: w = d.width
                    break
            if cd is not None and cd.width and w is None: w = cd.width
            colw.append(0 if hidden else col_px(w if w else dw))
        rowh = []
        for r in range(1, max(maxr, 1) + 1):
            rd = ws.row_dimensions.get(r)
            if rd is not None and rd.hidden: rowh.append(0)
            else: rowh.append(round(((rd.height if rd is not None and rd.height else dh)) * 96 / 72, 1))
        merged = {}
        skip = set()
        for mr in ws.merged_cells.ranges:
            if mr.min_row > maxr or mr.min_col > maxc: continue
            r2, c2 = min(mr.max_row, maxr), min(mr.max_col, maxc)
            merged[(mr.min_row, mr.min_col)] = (r2, c2)
            for r in range(mr.min_row, r2 + 1):
                for c in range(mr.min_col, c2 + 1):
                    if (r, c) != (mr.min_row, mr.min_col): skip.add((r, c))
        trs = []
        for r in range(1, maxr + 1):
            if rowh[r - 1] == 0: continue
            tds = []
            for c in range(1, maxc + 1):
                if (r, c) in skip or colw[c - 1] == 0:
                    if (r, c) not in skip and colw[c - 1] == 0: pass
                    continue
                cell = ws.cell(r, c)
                span = ''
                rr, cc2 = r, c
                if (r, c) in merged:
                    rr, cc2 = merged[(r, c)]
                    vis_cols = sum(1 for k in range(c, cc2 + 1) if colw[k - 1] > 0)
                    vis_rows = sum(1 for k in range(r, rr + 1) if rowh[k - 1] > 0)
                    if vis_cols > 1: span += ' colspan="%d"' % vis_cols
                    if vis_rows > 1: span += ' rowspan="%d"' % vis_rows
                st = ''
                if cell.has_style:
                    f = cell.font
                    if f is not None:
                        if f.b: st += 'font-weight:700;'
                        if f.i: st += 'font-style:italic;'
                        if f.u: st += 'text-decoration:underline;'
                        if f.sz and abs(float(f.sz) - 11) > 0.1: st += 'font-size:%.1fpt;' % float(f.sz)
                        fc = theme.xl(f.color) if f.color is not None else None
                        if fc and fc.upper() != '#000000': st += 'color:%s;' % fc
                    fl = cell.fill
                    if fl is not None and fl.fill_type == 'solid':
                        bg = theme.xl(fl.fgColor) or theme.xl(fl.start_color)
                        if bg: st += 'background:%s;' % bg
                    b = cell.border
                    br = ws.cell(r, cc2).border if cc2 != c else b
                    bb = ws.cell(rr, c).border if rr != r else b
                    if b is not None:
                        st += bcss('left', b.left, theme) + bcss('top', b.top, theme)
                    if br is not None: st += bcss('right', br.right, theme)
                    if bb is not None: st += bcss('bottom', bb.bottom, theme)
                    al = cell.alignment
                    if al is not None:
                        if al.horizontal in ('center', 'centerContinuous', 'distributed'): st += 'text-align:center;'
                        elif al.horizontal == 'right': st += 'text-align:right;'
                        elif al.horizontal == 'justify': st += 'text-align:justify;'
                        if al.vertical == 'top': st += 'vertical-align:top;'
                        elif al.vertical in ('center', 'distributed', 'justify'): st += 'vertical-align:middle;'
                        if al.wrap_text or al.shrink_to_fit: st += 'white-space:pre-wrap;overflow:hidden;'
                        if al.text_rotation == 255: st += 'writing-mode:vertical-rl;'
                        if al.indent: st += 'padding-left:%dpx;' % int(al.indent * 9)
                v = cell.value
                txt = '' if v is None else num_fmt(v, cell.number_format or 'General')
                if isinstance(v, (int, float)) and not isinstance(v, bool) and 'text-align' not in st: st += 'text-align:right;'
                tds.append('<td%s%s>%s</td>' % (span, ' style="%s"' % st if st else '', esc(txt).replace('\n', '<br>')))
            trs.append('<tr style="height:%.1fpx">%s</tr>' % (rowh[r - 1], ''.join(tds)))
        cols = ''.join('<col style="width:%dpx">' % w for w in colw if w > 0)
        tw = sum(colw)
        # 도형 레이어: 좌표는 열 누적(숨김 제외), 행은 JS 로 실제 높이에 맞춤
        colx = [0]
        for w in colw: colx.append(colx[-1] + w)
        ov = []
        for k, fr, to, el in anc:
            def cx(col, off):
                if col < len(colx): return colx[col] + off
                return colx[-1] + (col - len(colx) + 1) * col_px(dw) + off
            xf = xfrm_of(el) if loc(el.tag) != 'AlternateContent' else None
            if k == 'two':
                x1, x2 = cx(fr[0], fr[1]), cx(to[0], to[1])
                a = 'r:%d,%.1f,%d,%.1f' % (fr[2], fr[3], to[2], to[3])
                w0 = max(1, x2 - x1)
                # 기준 높이는 엑셀 행 높이(px)
                def ry(row, off):
                    return sum(rowh[:row]) + off if row <= len(rowh) else sum(rowh) + (row - len(rowh)) * dh * 96 / 72 + off
                h0 = max(1, ry(to[2], to[3]) - ry(fr[2], fr[3]))
            elif k == 'one':
                x1 = cx(fr[0], fr[1]); w0, h0 = to
                a = 'o:%d,%.1f,%.1f' % (fr[2], fr[3], h0)
            else:
                x1 = fr[0]; w0, h0 = to; a = 'a:%.1f,%.1f' % (fr[1], h0)
            if '터틀' in ws.title and loc(el.tag) == 'grpSp' and TURTLE_CUR.get('t'):
                w0 = max(400, min(w0, sum(colw) - x1 - 8)); inner = turtle_html(TURTLE_CUR['t'], w0, h0)
            elif '터틀' in ws.title and TURTLE_CUR.get('t') and loc(el.tag) == 'sp':
                continue
            else:
                inner = render_el_box(ctx, el, (0, 0, w0, h0), xf)
            ov.append('<div class="anc" data-a="%s" style="left:%.1fpx;width:%.1fpx;height:%.1fpx">%s</div>' % (a, x1, w0, h0, inner))
        sheets.append({'name': ws.title, 'rh': rowh, 'html': '<div class="sheet"><table class="xl" style="width:%dpx"><colgroup>%s</colgroup>%s</table><div class="ovl">%s</div></div>' % (tw, cols, ''.join(trs), ''.join(ov))})
    return sheets

# ───────────── PPTX ─────────────
def render_pptx(path):
    z = zipfile.ZipFile(path)
    pres = ET.fromstring(z.read('ppt/presentation.xml'))
    sz = kid(pres, 'sldSz'); SW, SH = int(sz.get('cx')), int(sz.get('cy'))
    pr = rels_of(z, 'ppt/presentation.xml')
    W = 960.0 if SW >= SH else 720.0
    sc = W / SW
    slides = []
    lst = kid(pres, 'sldIdLst')
    theme_part = None
    for n in z.namelist():
        if n.startswith('ppt/theme/theme') and n.endswith('.xml'): theme_part = theme_part or n
    theme = Theme(theme_colors(z, theme_part) if theme_part else {})
    def phinfo(part):
        """layout/master 의 placeholder (type, idx) → (xf, style)"""
        x = ET.fromstring(z.read(part)); out = {}
        for sp in x.iter():
            if loc(sp.tag) != 'sp': continue
            ph = None
            for e in sp.iter():
                if loc(e.tag) == 'ph': ph = e; break
            if ph is None: continue
            xf = xfrm_of(sp)
            bp = path_(sp, 'txBody', 'bodyPr')
            st = {'anchor': bp.get('anchor') if bp is not None else None}
            lv = path_(sp, 'txBody', 'lstStyle', 'lvl1pPr')
            if lv is not None:
                dr = kid(lv, 'defRPr')
                if dr is not None and dr.get('sz'): st['sz'] = int(dr.get('sz')) / 100
                if dr is not None and dr.get('b') == '1': st['b'] = True
                if lv.get('algn'): st['algn'] = lv.get('algn')
            out[(ph.get('type', 'body'), ph.get('idx'))] = (xf, st)
            out.setdefault(('type', ph.get('type', 'body')), (xf, st))
        return out
    def master_tx_sz(mpart):
        x = ET.fromstring(z.read(mpart)); r = {}
        for nm, key in (('titleStyle', 'title'), ('bodyStyle', 'body'), ('otherStyle', 'other')):
            e = None
            for c in x.iter():
                if loc(c.tag) == nm: e = c; break
            l1 = kid(e, 'lvl1pPr') if e is not None else None
            dr = kid(l1, 'defRPr') if l1 is not None else None
            if dr is not None and dr.get('sz'): r[key] = int(dr.get('sz')) / 100
        return r
    def bg_of(part):
        x = ET.fromstring(z.read(part))
        for e in x.iter():
            if loc(e.tag) == 'bgPr': return theme.fill_of(e)
            if loc(e.tag) == 'bgRef' and len(e): return theme.dml(e[0])
        return None
    def shapes_of(part, ctx, T, ph_lookup=None, skip_ph=False):
        x = ET.fromstring(z.read(part))
        tree = None
        for e in x.iter():
            if loc(e.tag) == 'spTree': tree = e; break
        out = []
        if tree is None: return ''
        for el in tree:
            n = loc(el.tag)
            if n not in ('sp', 'cxnSp', 'pic', 'grpSp', 'graphicFrame', 'AlternateContent'): continue
            if skip_ph:
                if any(loc(e.tag) == 'ph' for e in el.iter()): continue
            out.append(render_el(ctx, el, T, ph_lookup))
        return ''.join(out)
    T = (lambda x, y: (x * sc, y * sc), sc, sc)
    for sid in (kids(lst, 'sldId') if lst is not None else []):
        part = pr.get(sid.get('{%s}id' % R_NS))
        if not part: continue
        sr = rels_of(z, part)
        lay = next((v for v in sr.values() if 'slideLayouts/' in v), None)
        mas = next((v for v in rels_of(z, lay).values() if 'slideMasters/' in v), None) if lay else None
        lph = phinfo(lay) if lay else {}
        mph = phinfo(mas) if mas else {}
        msz = master_tx_sz(mas) if mas else {}
        def lookup(sp, lph=lph, mph=mph, msz=msz):
            ph = None
            for e in sp.iter():
                if loc(e.tag) == 'ph': ph = e; break
            if ph is None: return None, None
            t, i = ph.get('type', 'body'), ph.get('idx')
            for src in (lph, mph):
                v = src.get((t, i)) or (src.get(('type', t)) if t != 'body' or i is None else None)
                if v is None and i is not None:
                    v = next((vv for (tt, ii), vv in src.items() if ii == i), None)
                if v and v[0]:
                    st = dict(v[1]);
                    if 'sz' not in st:
                        st['sz'] = msz.get('title' if t in ('title', 'ctrTitle') else 'body', 18)
                    if t in ('title', 'ctrTitle'): st.setdefault('anchor', 'ctr')
                    return v[0], st
            return None, None
        bg = bg_of(part) or (bg_of(lay) if lay else None) or (bg_of(mas) if mas else None) or '#FFFFFF'
        body = ''
        Ctx.slide_no = len(slides) + 1
        if mas: body += shapes_of(mas, Ctx(z, mas, theme, 'pptx'), T, skip_ph=True)
        if lay: body += shapes_of(lay, Ctx(z, lay, theme, 'pptx'), T, skip_ph=True)
        body += shapes_of(part, Ctx(z, part, theme, 'pptx'), T, lookup)
        slides.append({'name': '%d' % (len(slides) + 1), 'html': '<div class="slide" style="width:%.0fpx;height:%.0fpx;background:%s">%s</div>' % (W, SH * sc, bg if bg != 'none' else '#fff', body)})
    return slides

# ───────────── 페이지 ─────────────
PAGE_CSS = '''
*{box-sizing:border-box}html,body{margin:0;background:#E9ECF1;font-family:"Malgun Gothic","맑은 고딕","Apple SD Gothic Neo","Noto Sans KR",sans-serif;color:#111}
.bar{position:sticky;top:0;z-index:5;display:flex;gap:6px;align-items:center;padding:8px 12px;background:#fff;border-bottom:1px solid #D5DAE2;flex-wrap:wrap;font-size:13px}
.bar .t{font-weight:700;margin-right:8px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;max-width:46vw}
.bar button{font:inherit;border:1px solid #C9D0DA;background:#F6F8FB;border-radius:6px;padding:4px 10px;cursor:pointer;color:#223}
.bar button.on{background:#1F3A8A;border-color:#1F3A8A;color:#fff}
.bar .sp{flex:1}.bar .z{color:#667;min-width:44px;text-align:center}
.wrap{padding:16px;overflow:auto}
.pg{display:none;margin:0 auto;width:max-content}.pg.on{display:block}
.mode-all .pg{display:block;margin-bottom:18px}
.sheet{position:relative;background:#fff;box-shadow:0 1px 4px rgba(0,0,0,.18);padding:0}
table.xl{border-collapse:collapse;table-layout:fixed;font-size:11pt;background:#fff}
table.xl td{padding:0 3px;overflow:visible;white-space:nowrap;vertical-align:bottom;line-height:1.25;font-size:inherit}
.ovl{position:absolute;left:0;top:0;width:0;height:0}
.anc{position:absolute}
.shp{position:absolute;overflow:visible}
.shp .geo{position:absolute;inset:0;width:100%;height:100%;overflow:visible}
.shp .tx{position:absolute;inset:0;display:flex;flex-direction:column;overflow:visible;line-height:1.15}
.shp p,.ptbl p{margin:0}
.ptbl{border-collapse:collapse;table-layout:fixed;font-size:10pt}
.ptbl td{border:1px solid #888;padding:2px 4px;vertical-align:top;line-height:1.2;overflow:hidden}
.turtle{position:absolute;left:0;top:0;display:grid;grid-template-columns:1fr 1.15fr 1fr;grid-template-rows:auto 1fr auto;gap:14px;padding:10px;font-size:10.5pt}
.turtle .tt{border:1.5px solid #333;background:#fff;display:flex;flex-direction:column}
.turtle .tt b{background:#FFCC99;border-bottom:1.5px solid #333;text-align:center;padding:5px;font-size:11pt}
.turtle ul{margin:0;padding:6px 8px 6px 22px;line-height:1.35}
.turtle .a{grid-area:1/1}.turtle .b{grid-area:1/3}.turtle .c{grid-area:2/1;align-self:center}.turtle .core{grid-area:2/2;position:relative;border:0;background:none;min-height:150px;align-items:center;justify-content:center}
.turtle .core svg{position:absolute;inset:0;width:100%;height:100%}.turtle .core span{position:relative;font-weight:800;font-size:17pt;text-align:center;line-height:1.15}
.turtle .d{grid-area:2/3;align-self:center}.turtle .e{grid-area:3/1}.turtle .f{grid-area:3/3}
.chart{border:1px dashed #aab;color:#889;display:flex;align-items:center;justify-content:center;font-size:12px}
.slide .tx{letter-spacing:-.03em}.slide{position:relative;overflow:hidden;box-shadow:0 1px 6px rgba(0,0,0,.25);font-size:18pt}
.snum{font-size:12px;color:#667;margin:0 0 4px}
@media print{.bar{display:none}.wrap{padding:0}.pg{display:block!important;page-break-after:always}.sheet,.slide{box-shadow:none}html,body{background:#fff}}
'''
PAGE_JS = r'''
(function(){
  var pgs=[].slice.call(document.querySelectorAll('.pg')),tabs=[].slice.call(document.querySelectorAll('[data-pg]'));
  var z=1,fit=true,all=document.body.classList.contains('mode-all');
  function cur(){return pgs.filter(function(p){return p.classList.contains('on')})[0]||pgs[0]}
  function place(sh){
    var tb=sh.querySelector('table.xl');if(!tb)return;
    var rows=tb.rows,ry=[];var y=0;var hs=JSON.parse(sh.getAttribute('data-rh')||'[]');
    var vis=[];for(var i=0;i<hs.length;i++){vis.push(hs[i]>0)}
    var k=0;for(var r=0;r<hs.length;r++){ry.push(y);if(vis[r]){var tr=rows[k++];var bh=tr?tr.getBoundingClientRect().height/(parseFloat(sh.closest('.pg').style.zoom)||1):0;y+=bh>0?bh:hs[r]}}
    ry.push(y);
    function Y(row,off){return row<ry.length?ry[row]+off:y+(row-ry.length+1)*20+off}
    [].forEach.call(sh.querySelectorAll('.anc'),function(a){
      var d=a.getAttribute('data-a'),t=d.slice(0,1),v=d.slice(2).split(',').map(Number);
      if(t==='r'){var y1=Y(v[0],v[1]),y2=Y(v[2],v[3]);a.style.top=y1+'px';var h0=parseFloat(a.style.height);var h=Math.max(1,y2-y1);
        if(Math.abs(h-h0)>1){var s=h/h0;[].forEach.call(a.children,function(c){c.style.transformOrigin='0 0';c.style.transform=(c.style.transform||'')+' scaleY('+s+')'});a.style.height=h+'px'}}
      else if(t==='o'){a.style.top=Y(v[0],v[1])+'px'}else{a.style.top=v[0]+'px'}
    });
  }
  function zoom(){var p=all?pgs:[cur()];var w=document.querySelector('.wrap').clientWidth-34;
    p.forEach(function(pg){var el=pg.querySelector('.sheet,.slide');if(!el)return;var zz=z;if(fit){var nw=el.scrollWidth;zz=Math.min(1.6,Math.max(.3,w/nw));if(el.classList.contains('slide')&&!all){var hh=window.innerHeight-document.querySelector('.bar').offsetHeight-60;zz=Math.max(.3,Math.min(zz,hh/el.scrollHeight));}}pg.style.zoom=zz;});
    document.querySelector('.z').textContent=fit?'맞춤':Math.round(z*100)+'%';}
  function placeAll(){pgs.forEach(function(p){if(p.offsetParent!==null){var s=p.querySelector('.sheet');if(s&&!s._pl){place(s);s._pl=1}}})}
  function show(i){pgs.forEach(function(p,j){p.classList.toggle('on',j===i)});tabs.forEach(function(t,j){t.classList.toggle('on',j===i)});zoom();placeAll();try{location.hash='p'+(i+1)}catch(e){}}
  tabs.forEach(function(t,i){t.onclick=function(){show(i)}});
  document.getElementById('zi').onclick=function(){fit=false;z=Math.min(3,(z||1)+.15);zoom()};
  document.getElementById('zo').onclick=function(){fit=false;z=Math.max(.3,(z||1)-.15);zoom()};
  document.getElementById('zf').onclick=function(){fit=!fit;if(!fit)z=1;zoom()};
  var pa=document.getElementById('pa');if(pa)pa.onclick=function(){all=!all;document.body.classList.toggle('mode-all',all);pa.classList.toggle('on',all);zoom();placeAll()};
  document.getElementById('pr').onclick=function(){window.print()};
  window.addEventListener('resize',zoom);
    var m=/p(\d+)/.exec(location.hash||''),q=/[?&]s=([^&]+)/.exec(location.search||'');var st=0;
  if(m)st=+m[1]-1;else if(q){var nm=decodeURIComponent(q[1]);tabs.forEach(function(t,i){if(t.textContent===nm)st=i})}
  show(Math.min(st,pgs.length-1));
})();
'''

def page(title, items, kind):
    tabs = ''.join('<button data-pg="%d">%s</button>' % (i, esc(it['name'])) for i, it in enumerate(items)) if len(items) > 1 else ''
    if kind == 'pptx' and len(items) > 12: tabs = ''.join('<button data-pg="%d">%s</button>' % (i, i + 1) for i, it in enumerate(items))
    pgs = []
    for i, it in enumerate(items):
        h = it['html']
        if kind == 'xlsx':
            h = h.replace('<div class="sheet">', '<div class="sheet" data-rh="%s">' % esc(json.dumps(it.get('rh', []))), 1)
        pgs.append('<div class="pg">%s%s</div>' % ('<div class="snum">%s</div>' % (('슬라이드 %d' % (i + 1)) if kind == 'pptx' else '') if kind == 'pptx' else '', h))
    allbtn = '<button id="pa">전체 펼치기</button>' if len(items) > 1 else ''
    return '''<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>%s</title><style>%s</style></head>
<body><div class="bar"><span class="t">%s</span>%s<span class="sp"></span>%s<button id="zo">－</button><span class="z">맞춤</span><button id="zi">＋</button><button id="zf">맞춤/원래 크기</button><button id="pr">인쇄</button></div>
<div class="wrap">%s</div><script>%s</script></body></html>''' % (esc(title), PAGE_CSS, esc(title), tabs, allbtn, ''.join(pgs), PAGE_JS)

# ───────────── 터틀맵 ─────────────
TURTLE_KEYS = [('measure', ('측정',)), ('input', ('입력',)), ('output', ('출력',)),
               ('what', ('물적자원', '설비/재료', '설비', '재료')), ('how', ('절차', '방법', '규정')), ('who', ('인적자원', '사람'))]

def turtle_of(path):
    z = zipfile.ZipFile(path)
    parts = sheet_parts(z)
    dpart = None
    for nm, (p, d) in parts.items():
        if '터틀' in nm and d: dpart = d
    if not dpart: return None
    x = ET.fromstring(z.read(dpart))
    items = []
    for sp in x.iter():
        if loc(sp.tag) != 'sp': continue
        paras = []
        for p in sp.iter():
            if loc(p.tag) != 'p': continue
            t = ''.join((r.text or '') for r in p.iter() if loc(r.tag) == 't')
            t = re.sub(r'\s+', ' ', t).strip()
            if t: paras.append(t)
        if paras: items.append(paras)
    if not items: return None
    res = {'name': re.sub(r'\s*PROCESS\s*$', '', ' '.join(items[0])).replace(' 관 리', '관리').strip()}
    i = 1
    while i < len(items) - 1:
        label = ' '.join(items[i])
        key = None
        for k, words in TURTLE_KEYS:
            if any(label.startswith(w) for w in words) and len(label) <= 12: key = k; break
        if key:
            vals = []
            for t in items[i + 1]:
                t = t.lstrip('●').strip()
                if not t: continue
                # 이어지는 줄(들여쓴 보조 줄) 병합
                if vals and (t.startswith('(') or t.startswith('/') or t.startswith('E-mail') or t.startswith('계획서') or t.startswith('등에 대한')):
                    vals[-1] = (vals[-1].rstrip(', ') + (' ' if not t.startswith('/') else '') + t).strip(); continue
                for part in re.split(r'\s{3,}●?\s*|\s*●\s*', t):
                    part = re.sub(r'^\d\.\s*', '', part).strip()
                    if part: vals.append(part)
            res[key] = vals; res[key + 'Label'] = label; i += 2
        else: i += 1
    return res

# ───────────── 메인 ─────────────
def main():
    import sys
    flt = [nfc(a) for a in sys.argv[1:]]
    os.makedirs(IMG, exist_ok=True)
    if not flt:
        for f in glob.glob(os.path.join(OUT, '*.html')): os.remove(f)
    files = []
    for d, _, fs in os.walk(SRC):
        if '_logs' in d: continue
        for f in fs:
            if f.lower().endswith(('.xlsx', '.pptx')) and not f.startswith('~$'): files.append(os.path.join(d, f))
    files.sort()
    index, turtles, fails = {}, {}, []
    if flt:
        files = [p for p in files if any(a in nfc(p) for a in flt)]
        print('partial', len(files)); 

    for p in files:
        rel = '../' + os.path.relpath(p, ROOT).replace(os.sep, '/')
        rel = nfc(rel)
        title = nfc(os.path.splitext(os.path.basename(p))[0])
        h = hashlib.sha1(rel.encode('utf-8')).hexdigest()[:12]
        try:
            if p.lower().endswith('.xlsx'):
                items = render_xlsx(p); kind = 'xlsx'
                # 행 높이 정보 → JS 재배치용
                wb = None
            else:
                items = render_pptx(p); kind = 'pptx'
        except Exception as e:
            import traceback; traceback.print_exc()
            fails.append((title, repr(e))); continue
        if not items: fails.append((title, 'empty')); continue
        for it in items:
            if kind == 'xlsx':
                # data-rh: 행별 기준 높이(숨김 0) — 표 렌더시 사용한 값과 동일하게 재계산
                pass
        html_ = page(title, items, kind)
        with open(os.path.join(OUT, h + '.html'), 'w', encoding='utf-8') as f: f.write(html_)
        index[rel] = {'p': 'preview/' + h + '.html', 'k': kind, 's': [it['name'] for it in items]}
        if re.match(r'^MP-\d{4}', title) and kind == 'xlsx':
            t = turtle_of(p)
            if t: turtles[title[:7]] = t
    if flt:
        print('partial build — index not rewritten', [index[k]['p'] for k in index]); return
    with open(os.path.join(ROOT, 'app', 'data', 'previews.js'), 'w', encoding='utf-8') as f:
        f.write('/* 자동 생성: scripts/build_preview.py — 원본(정정본) 경로 → 화면 미리보기 */\nwindow.PREVIEWS = ' + json.dumps(index, ensure_ascii=False, indent=0) + ';\n')
    with open(os.path.join(ROOT, 'app', 'data', 'turtle.js'), 'w', encoding='utf-8') as f:
        f.write('/* 자동 생성: scripts/build_preview.py — 프로세스 원본 \'터틀분석\' 시트 */\nwindow.TURTLE = ' + json.dumps(turtles, ensure_ascii=False, indent=1) + ';\n')
    print('previews', len(index), 'turtles', len(turtles), 'fails', fails)

if __name__ == '__main__':
    main()
